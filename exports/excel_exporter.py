"""
JurisIntel — Exportador para Excel (.xlsx).

Gera planilha profissional com:
  - Aba "Decisões": todos os campos com auto-filtro e primeira linha congelada
  - Aba "Resumo": tabelas de contagem por área, resultado, tipo de ação, relator
  - Formatação TJ-SP (cabeçalho azul escuro, fontes Arial, bordas finas)
"""

import os
import json
import logging
import argparse
from datetime import datetime

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

import sys
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from database.db import JurisprudenciaDB

logger = logging.getLogger(__name__)

# ── Estilos ───────────────────────────────────────────────────────────────────
_HDR_FILL   = PatternFill('solid', fgColor='1F4E79')
_HDR_FONT   = Font(name='Arial', bold=True, color='FFFFFF', size=10)
_CELL_FONT  = Font(name='Arial', size=9)
_TITLE_FONT = Font(name='Arial', bold=True, size=13, color='1F4E79')
_THIN_BORDER = Border(
    left=Side(style='thin', color='D9D9D9'),
    right=Side(style='thin', color='D9D9D9'),
    top=Side(style='thin', color='D9D9D9'),
    bottom=Side(style='thin', color='D9D9D9'),
)

COLUNAS = {
    'numero_processo':    ('Nº Processo',          25),
    'classe_processual':  ('Classe',                18),
    'area_direito':       ('Área do Direito',        22),
    'tipo_acao':          ('Tipo de Ação',           24),
    'tese_principal':     ('Tese Principal',         32),
    'tese_secundaria':    ('Tese Secundária',         26),
    'resultado':          ('Resultado',              15),
    'resultado_detalhado':('Resultado Detalhado',    36),
    'autor_tipo':         ('Autor (Tipo)',           15),
    'autor_nome':         ('Autor',                  26),
    'reu_tipo':           ('Réu (Tipo)',             20),
    'reu_nome':           ('Réu',                    26),
    'comarca':            ('Comarca',                15),
    'relator':            ('Relator',                28),
    'orgao_julgador':     ('Órgão Julgador',         28),
    'data_julgamento':    ('Data Julgamento',        14),
    'data_publicacao':    ('Data Publicação',        14),
    'valor_causa':        ('Valor da Causa (R$)',    16),
    'valor_condenacao':   ('Valor Condenação (R$)',  16),
    'dano_moral_valor':   ('Dano Moral (R$)',        16),
    'fundamentacao_legal':('Fundamentação Legal',   32),
    'precedentes_citados':('Precedentes',            26),
    'palavras_chave':     ('Palavras-chave',         32),
    'resumo_ia':          ('Resumo IA',              52),
    'ementa':             ('Ementa',                 52),
}


def _fmt_json(val):
    if not val:
        return ''
    if isinstance(val, str):
        try:
            parsed = json.loads(val)
            if isinstance(parsed, list):
                return '; '.join(str(x) for x in parsed)
        except (json.JSONDecodeError, TypeError):
            pass
        return val
    if isinstance(val, list):
        return '; '.join(str(x) for x in val)
    return str(val)


def exportar_excel(decisoes: list, output_path: str,
                   titulo: str = 'Jurisprudência TJ-SP'):
    """
    Exporta lista de decisões para Excel formatado.

    Args:
        decisoes:    lista de dicts (do banco)
        output_path: caminho de saída (.xlsx)
        titulo:      título da planilha
    """
    if not decisoes:
        logger.warning("Nenhuma decisão para exportar.")
        return

    os.makedirs(os.path.dirname(output_path) or '.', exist_ok=True)

    json_fields = {'fundamentacao_legal', 'precedentes_citados', 'palavras_chave'}
    dados = []
    for d in decisoes:
        row = {}
        for campo, (nome, _) in COLUNAS.items():
            val = d.get(campo, '')
            row[nome] = _fmt_json(val) if campo in json_fields else (val or '')
        dados.append(row)

    df = pd.DataFrame(dados)
    wb = Workbook()

    # ── Aba 1: Decisões ───────────────────────────────────────────────────────
    ws = wb.active
    ws.title = 'Decisões'

    for col_idx, col_name in enumerate(df.columns, 1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.font      = _HDR_FONT
        cell.fill      = _HDR_FILL
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border    = _THIN_BORDER
        largura = list(COLUNAS.values())[col_idx - 1][1]
        ws.column_dimensions[get_column_letter(col_idx)].width = largura

    for row_idx, row_data in enumerate(df.itertuples(index=False), 2):
        for col_idx, val in enumerate(row_data, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.font      = _CELL_FONT
            cell.border    = _THIN_BORDER
            cell.alignment = Alignment(vertical='top', wrap_text=True)

    ws.auto_filter.ref = ws.dimensions
    ws.freeze_panes    = 'A2'
    ws.row_dimensions[1].height = 30

    # ── Aba 2: Resumo ─────────────────────────────────────────────────────────
    ws2 = wb.create_sheet('Resumo')
    ws2.sheet_properties.tabColor = '1F4E79'
    ws2['A1'] = titulo
    ws2['A1'].font = _TITLE_FONT
    ws2.merge_cells('A1:D1')
    ws2['A2'] = f"Gerado em: {datetime.now():%d/%m/%Y %H:%M}  —  {len(decisoes)} decisões"
    ws2['A2'].font = Font(name='Arial', size=9, italic=True, color='888888')
    ws2.column_dimensions['A'].width = 42
    ws2.column_dimensions['B'].width = 14

    resumos = [
        ('Área do Direito',  'area_direito'),
        ('Resultado',        'resultado'),
        ('Tipo de Ação',     'tipo_acao'),
        ('Tese Principal',   'tese_principal'),
        ('Tipo do Réu',      'reu_tipo'),
        ('Relator',          'relator'),
        ('Comarca',          'comarca'),
    ]
    row_n = 4
    for titulo_tab, campo in resumos:
        # cabeçalho da tabela
        c = ws2.cell(row=row_n, column=1, value=titulo_tab)
        c.font = Font(name='Arial', bold=True, size=10, color='1F4E79')
        ws2.cell(row=row_n, column=2, value='Qtd').font = \
            Font(name='Arial', bold=True, size=10, color='1F4E79')
        row_n += 1

        contagem: dict = {}
        for d in decisoes:
            v = d.get(campo) or 'Não classificado'
            contagem[v] = contagem.get(v, 0) + 1

        for v, cnt in sorted(contagem.items(), key=lambda x: -x[1])[:20]:
            ws2.cell(row=row_n, column=1, value=v).font   = _CELL_FONT
            ws2.cell(row=row_n, column=2, value=cnt).font = _CELL_FONT
            row_n += 1
        row_n += 1

    wb.save(output_path)
    logger.info(f"Excel exportado: {output_path} ({len(decisoes)} decisões)")


def main():
    logging.basicConfig(level=logging.INFO)
    parser = argparse.ArgumentParser(description='Exportador Excel — JurisIntel')
    parser.add_argument('--todos',     action='store_true')
    parser.add_argument('--area',      help='Filtrar por área')
    parser.add_argument('--tipo-acao', help='Filtrar por tipo de ação')
    parser.add_argument('--tese',      help='Filtrar por tese (parcial)')
    parser.add_argument('--resultado', help='Filtrar por resultado')
    parser.add_argument('--reu-tipo',  help='Filtrar por tipo de réu')
    parser.add_argument('--comarca')
    parser.add_argument('--periodo',   help='Ano (ex: 2025)')
    parser.add_argument('--busca-fts', help='Busca full-text')
    parser.add_argument('--output', '-o')
    args = parser.parse_args()

    db = JurisprudenciaDB()

    if args.busca_fts:
        decisoes = db.buscar_fts(args.busca_fts, limit=5000)
    else:
        filtros = {}
        if args.area:      filtros['area_direito']  = args.area
        if args.tipo_acao: filtros['tipo_acao']     = args.tipo_acao
        if args.tese:      filtros['tese_principal'] = args.tese
        if args.resultado: filtros['resultado']     = args.resultado
        if args.reu_tipo:  filtros['reu_tipo']      = args.reu_tipo
        if args.comarca:   filtros['comarca']       = args.comarca
        if args.periodo:
            filtros['data_inicio'] = f'{args.periodo}-01-01'
            filtros['data_fim']    = f'{args.periodo}-12-31'
        decisoes = db.consultar(filtros=filtros if filtros else None, limit=50000)

    if not decisoes:
        print("Nenhuma decisão encontrada.")
        return

    ts      = datetime.now().strftime('%Y%m%d_%H%M%S')
    outpath = args.output or f'data/exports/jurisprudencia_{ts}.xlsx'
    titulo_partes = ['Jurisprudência TJ-SP']
    if args.area:  titulo_partes.append(f"Área: {args.area}")
    if args.tese:  titulo_partes.append(f"Tese: {args.tese}")

    exportar_excel(decisoes, outpath, titulo=' | '.join(titulo_partes))
    print(f"\n✓ Exportado: {outpath} ({len(decisoes)} decisões)")


if __name__ == '__main__':
    main()
